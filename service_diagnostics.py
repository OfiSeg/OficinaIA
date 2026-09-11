"""Autodiagnóstico seguro de dependencias reales de OficinaIA.

Los checks son de sólo lectura. Ninguno escribe datos, envía correos, modifica
Sheets/Excel ni ejecuta análisis de documentos. Los resultados nunca contienen
secretos.
"""
from __future__ import annotations

import re
import threading
import time
from dataclasses import dataclass, asdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Callable

import runtime_config

_CACHE_LOCK = threading.Lock()
_CACHE: dict[str, object] = {"at": 0.0, "result": None, "external": False}
CACHE_SECONDS = 30.0


@dataclass
class ServiceStatus:
    service: str
    label: str
    status: str  # ok | warning | error
    code: str
    message: str
    checks: list[dict]
    checked_at: str

    def as_dict(self) -> dict:
        return asdict(self)


def _now_iso() -> str:
    return datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")


def _check(label: str, ok: bool | None, detail: str) -> dict:
    # ok=None significa advertencia/estado informativo.
    return {"label": label, "ok": ok, "detail": detail}


def _exc_code(exc: Exception, default: str) -> str:
    return str(getattr(exc, "code", "") or default)


def _safe_error_text(exc: Exception) -> str:
    """Mensaje técnico breve con redacción defensiva de secretos e IDs."""
    text = str(exc or "").strip().replace("\n", " ")
    if not text:
        return "Error no especificado."
    lowered = text.lower()
    if any(token in lowered for token in ("private_key", "client_secret", "refresh_token", "password")):
        return "El proveedor rechazó la configuración o las credenciales."
    text = re.sub(r"(?i)postgres(?:ql)?://[^\s]+", "DATABASE_URL=[OCULTO]", text)
    text = re.sub(r"(?i)(api[_-]?key|token|secret)\s*[:=]\s*[^\s,;]+", r"\1=[OCULTO]", text)
    text = re.sub(r"(?i)(spreadsheetId=)[A-Za-z0-9_-]+", r"\1[OCULTO]", text)
    text = re.sub(r"(?i)(/spreadsheets/)[A-Za-z0-9_-]+", r"\1[OCULTO]", text)
    if len(text) > 240:
        text = text[:237] + "..."
    return text


def check_google_sheets(*, probe_external: bool) -> ServiceStatus:
    import google_sheets_service as sheets

    checks: list[dict] = []
    sid = runtime_config.inspect_env("SHEET_ID_ASEGURADOS")
    creds = runtime_config.inspect_env("GOOGLE_SHEETS_CREDENTIALS_JSON")
    checks.append(_check("SHEET_ID", sid.configured, "Detectado" if sid.configured else ("Vacío" if sid.state == "empty" else "No configurado")))
    checks.append(_check("Credenciales", creds.configured, "Detectadas" if creds.configured else ("Vacías" if creds.state == "empty" else "No configuradas")))

    if not sid.configured:
        code = "EMPTY_SHEET_ID" if sid.state == "empty" else "MISSING_SHEET_ID"
        return ServiceStatus("google_sheets", "Google Sheets", "error", code, "Falta configurar la planilla de asegurados.", checks, _now_iso())
    if not creds.configured:
        code = "EMPTY_GOOGLE_CREDENTIALS" if creds.state == "empty" else "MISSING_GOOGLE_CREDENTIALS"
        return ServiceStatus("google_sheets", "Google Sheets", "error", code, "Faltan las credenciales de Google Sheets.", checks, _now_iso())

    try:
        # Parsear credenciales es local y permite diferenciar JSON inválido sin red.
        sheets.validate_credentials_configuration()
        checks.append(_check("Formato credenciales", True, "Válido"))
    except Exception as exc:
        checks.append(_check("Formato credenciales", False, "Inválido"))
        return ServiceStatus("google_sheets", "Google Sheets", "error", _exc_code(exc, "INVALID_GOOGLE_CREDENTIALS"), _safe_error_text(exc), checks, _now_iso())

    if not probe_external:
        checks.append(_check("Conexión", None, "No comprobada todavía"))
        return ServiceStatus("google_sheets", "Google Sheets", "warning", "NOT_PROBED", "Configuración detectada. Tocá “Comprobar servicios” para verificar acceso real.", checks, _now_iso())

    try:
        result = sheets.check_connection("1")
        checks.extend(result.get("checks") or [])
        return ServiceStatus("google_sheets", "Google Sheets", "ok", "OK", "Google Sheets conectado correctamente.", checks, _now_iso())
    except Exception as exc:
        code = _exc_code(exc, "GOOGLE_API_ERROR")
        checks.append(_check("Conexión", False, _safe_error_text(exc)))
        status = "warning" if code in {"GOOGLE_API_TEMPORARY", "GOOGLE_API_RATE_LIMIT", "SHEET_TIMEOUT"} else "error"
        return ServiceStatus("google_sheets", "Google Sheets", status, code, _safe_error_text(exc), checks, _now_iso())


def _classify_gemini_error(exc: Exception) -> tuple[str, str]:
    text = str(exc or "").upper()
    status_code = getattr(exc, "status_code", None) or getattr(exc, "code", None)
    if status_code == 401 or "401" in text or "UNAUTHENTICATED" in text or "API_KEY_INVALID" in text:
        return "GEMINI_INVALID_KEY", "La clave de Gemini fue rechazada."
    if status_code == 429 or "429" in text or "RESOURCE_EXHAUSTED" in text:
        return "GEMINI_QUOTA_EXCEEDED", "Gemini está temporalmente limitado por cuota."
    if status_code == 404 or "404" in text or "NOT_FOUND" in text:
        return "GEMINI_MODEL_UNAVAILABLE", "El modelo consultado no está disponible."
    if status_code in {500, 502, 503, 504} or any(x in text for x in ("500", "502", "503", "504", "UNAVAILABLE")):
        return "GEMINI_TEMPORARILY_UNAVAILABLE", "Gemini está temporalmente no disponible."
    if "TIMEOUT" in text or "TIMED OUT" in text or isinstance(exc, TimeoutError):
        return "GEMINI_TIMEOUT", "Gemini no respondió a tiempo."
    return "GEMINI_PROVIDER_ERROR", "No se pudo comprobar Gemini."


def check_gemini(*, probe_external: bool) -> ServiceStatus:
    key = runtime_config.inspect_env("GEMINI_API_KEY")
    checks = [_check("Configuración", key.configured, "Clave detectada" if key.configured else ("Clave vacía" if key.state == "empty" else "Clave no configurada"))]
    if not key.configured:
        code = "GEMINI_EMPTY_KEY" if key.state == "empty" else "GEMINI_MISSING_KEY"
        return ServiceStatus("gemini", "Gemini", "error", code, "Gemini no está configurado.", checks, _now_iso())
    try:
        from ai_gateway import obtener_cliente_gemini, DEFAULT_MODELS
        client = obtener_cliente_gemini()
        if client is None:
            raise RuntimeError("GEMINI_API_KEY no disponible")
        checks.append(_check("Cliente", True, "Inicializado"))
        if not probe_external:
            checks.append(_check("Proveedor", None, "No comprobado todavía"))
            return ServiceStatus("gemini", "Gemini", "warning", "NOT_PROBED", "Configuración detectada. Tocá “Comprobar servicios” para verificar el proveedor.", checks, _now_iso())
        # GET de metadatos: autentica sin generar contenido. Se prueban los
        # modelos reales del gateway para no declarar caída total si sólo uno
        # de los fallbacks dejó de estar disponible.
        last_exc = None
        model_unavailable = 0
        for model in DEFAULT_MODELS:
            try:
                client.models.get(model=model)
                checks.append(_check("Proveedor", True, "Accesible"))
                return ServiceStatus("gemini", "Gemini", "ok", "OK", "Gemini operativo.", checks, _now_iso())
            except Exception as exc:
                last_exc = exc
                code, _message = _classify_gemini_error(exc)
                if code == "GEMINI_MODEL_UNAVAILABLE":
                    model_unavailable += 1
                    continue
                raise
        if model_unavailable:
            checks.append(_check("Proveedor", False, "Los modelos configurados no están disponibles"))
            return ServiceStatus("gemini", "Gemini", "warning", "GEMINI_MODEL_UNAVAILABLE", "Gemini responde, pero los modelos configurados no están disponibles.", checks, _now_iso())
        if last_exc is not None:
            raise last_exc
        raise RuntimeError("No hay modelos Gemini configurados.")
    except Exception as exc:
        code, message = _classify_gemini_error(exc)
        checks.append(_check("Proveedor", False, message))
        status = "warning" if code in {"GEMINI_QUOTA_EXCEEDED", "GEMINI_TIMEOUT", "GEMINI_TEMPORARILY_UNAVAILABLE", "GEMINI_MODEL_UNAVAILABLE"} else "error"
        return ServiceStatus("gemini", "Gemini", status, code, message, checks, _now_iso())


def check_mail(*, probe_external: bool) -> ServiceStatus:
    names = ("GMAIL_SENDER_EMAIL", "GMAIL_OAUTH_CLIENT_ID", "GMAIL_OAUTH_CLIENT_SECRET", "GMAIL_OAUTH_REFRESH_TOKEN")
    states = [runtime_config.inspect_env(n) for n in names]
    configured_all = all(s.configured for s in states)
    checks = [_check("Configuración", configured_all, "Completa" if configured_all else "Incompleta")]
    if not configured_all:
        return ServiceStatus("mail", "Correo", "error", "MAIL_CONFIGURATION_ERROR", "Falta completar la configuración de correo.", checks, _now_iso())
    try:
        from mail_service import MailChannel
        _sender, creds = MailChannel._credentials()
        checks.append(_check("Credenciales", True, "Detectadas"))
        if not probe_external:
            checks.append(_check("Autenticación", None, "No comprobada todavía"))
            return ServiceStatus("mail", "Correo", "warning", "NOT_PROBED", "Configuración detectada. La comprobación no envía ningún correo.", checks, _now_iso())
        from google.auth.transport.requests import Request as GoogleAuthRequest
        creds.refresh(GoogleAuthRequest())
        checks.append(_check("Autenticación", True, "Correcta"))
        return ServiceStatus("mail", "Correo", "ok", "OK", "Correo operativo.", checks, _now_iso())
    except Exception as exc:
        text = str(exc or "").upper()
        if any(x in text for x in ("INVALID_GRANT", "UNAUTHORIZED", "401", "INVALID_CLIENT")):
            code, message = "MAIL_AUTH_ERROR", "Google rechazó las credenciales de correo."
        elif "TIMEOUT" in text or "CONNECTION" in text:
            code, message = "MAIL_CONNECTION_ERROR", "No se pudo conectar al servicio de correo."
        else:
            code, message = "MAIL_CONFIGURATION_ERROR", _safe_error_text(exc)
        checks.append(_check("Autenticación", False, message))
        return ServiceStatus("mail", "Correo", "error", code, message, checks, _now_iso())


def check_database(*, probe_external: bool) -> ServiceStatus:
    db_state = runtime_config.inspect_env("DATABASE_URL")
    checks: list[dict] = []
    if db_state.configured:
        checks.append(_check("Configuración", True, "PostgreSQL configurado"))
        if not probe_external:
            checks.append(_check("Conexión", None, "No comprobada todavía"))
            return ServiceStatus("database", "Base de datos", "warning", "NOT_PROBED", "PostgreSQL configurado. Tocá “Comprobar servicios” para probar la conexión.", checks, _now_iso())
        try:
            from database_pg import conectar_pg
            conn = conectar_pg()
            try:
                with conn.cursor() as cur:
                    cur.execute("SELECT 1")
                    cur.fetchone()
            finally:
                conn.close()
            checks.append(_check("Conexión", True, "Correcta"))
            return ServiceStatus("database", "Base de datos", "ok", "OK", "Base de datos operativa.", checks, _now_iso())
        except Exception as exc:
            checks.append(_check("Conexión", False, "Falló"))
            return ServiceStatus("database", "Base de datos", "error", "DATABASE_CONNECTION_ERROR", "No se pudo conectar a la base de datos.", checks, _now_iso())

    # Sin DATABASE_URL el proyecto usa SQLite local deliberadamente.
    checks.append(_check("Modo", True, "SQLite local"))
    try:
        from local_db import conectar_db
        conn = conectar_db()
        try:
            conn.execute("SELECT 1").fetchone()
        finally:
            conn.close()
        checks.append(_check("Conexión", True, "Correcta"))
        return ServiceStatus("database", "Base de datos", "ok", "SQLITE_LOCAL", "SQLite local operativo.", checks, _now_iso())
    except Exception as exc:
        checks.append(_check("Conexión", False, "Falló"))
        return ServiceStatus("database", "Base de datos", "error", "DATABASE_CONNECTION_ERROR", "No se pudo conectar a la base de datos.", checks, _now_iso())


def check_excel_local(*, probe_external: bool) -> ServiceStatus:
    path = Path(__file__).resolve().parent / "excel_interno.xlsx"
    checks = [_check("Archivo", path.is_file(), "Disponible" if path.is_file() else "No encontrado")]
    if not path.is_file():
        return ServiceStatus("excel_local", "Excel local", "warning", "EXCEL_NOT_FOUND", "El Excel local de referencia no está disponible.", checks, _now_iso())
    try:
        # El Excel local no es la fuente activa de Asegurados; se prueba en modo
        # read_only y sólo su estructura mínima.
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            if not wb.sheetnames:
                raise RuntimeError("El archivo no contiene hojas.")
            ws = wb[wb.sheetnames[0]]
            next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        finally:
            wb.close()
        checks.append(_check("Lectura", True, "Correcta"))
        return ServiceStatus("excel_local", "Excel local", "ok", "OK", "Excel local accesible (referencia/respaldo; la fuente activa es Google Sheets).", checks, _now_iso())
    except Exception as exc:
        checks.append(_check("Lectura", False, "Falló"))
        return ServiceStatus("excel_local", "Excel local", "warning", "EXCEL_READ_ERROR", _safe_error_text(exc), checks, _now_iso())


def check_storage(*, probe_external: bool) -> ServiceStatus:
    names = ("R2_ENDPOINT_URL", "R2_ACCESS_KEY_ID", "R2_SECRET_ACCESS_KEY", "R2_BUCKET_NAME")
    configured_all = runtime_config.configured(names)
    checks = [_check("Configuración", configured_all, "Completa" if configured_all else "Incompleta")]
    if not configured_all:
        return ServiceStatus("storage", "Almacenamiento", "warning", "STORAGE_CONFIGURATION_ERROR", "Cloudflare R2 no está completamente configurado.", checks, _now_iso())
    if not probe_external:
        checks.append(_check("Conexión", None, "No comprobada todavía"))
        return ServiceStatus("storage", "Almacenamiento", "warning", "NOT_PROBED", "R2 configurado. Tocá “Comprobar servicios” para verificar acceso real.", checks, _now_iso())
    try:
        import storage_r2
        storage_r2.comprobar_conexion()
        checks.append(_check("Conexión", True, "Correcta"))
        return ServiceStatus("storage", "Almacenamiento", "ok", "OK", "Almacenamiento operativo.", checks, _now_iso())
    except Exception as exc:
        checks.append(_check("Conexión", False, "Falló"))
        return ServiceStatus("storage", "Almacenamiento", "warning", "STORAGE_CONNECTION_ERROR", "No se pudo comprobar el almacenamiento.", checks, _now_iso())


CHECKERS: tuple[Callable[..., ServiceStatus], ...] = (
    check_google_sheets,
    check_gemini,
    check_mail,
    check_excel_local,
    check_database,
    check_storage,
)


def run_health_checks(*, probe_external: bool = False, use_cache: bool = True) -> dict:
    now = time.monotonic()
    with _CACHE_LOCK:
        cached = _CACHE.get("result")
        cached_external = bool(_CACHE.get("external"))
        if use_cache and cached and now - float(_CACHE.get("at") or 0.0) < CACHE_SECONDS:
            # Una comprobación externa satisface también un pedido liviano; una
            # comprobación local NO satisface un refresh explícito externo.
            if cached_external or not probe_external:
                return dict(cached)

    services: dict[str, dict] = {}
    for checker in CHECKERS:
        try:
            status = checker(probe_external=probe_external)
        except Exception as exc:  # el panel de salud jamás debe caerse solo
            name = getattr(checker, "__name__", "servicio").removeprefix("check_")
            status = ServiceStatus(name, name.replace("_", " ").title(), "error", "HEALTH_CHECK_INTERNAL_ERROR", _safe_error_text(exc), [], _now_iso())
        services[status.service] = status.as_dict()

    levels = [s["status"] for s in services.values()]
    overall = "error" if "error" in levels else ("warning" if "warning" in levels else "ok")
    result = {
        "status": overall,
        "services": services,
        "checked_at": _now_iso(),
        "external_probe": bool(probe_external),
        "cache_seconds": int(CACHE_SECONDS),
    }
    with _CACHE_LOCK:
        _CACHE.update({"at": now, "result": dict(result), "external": bool(probe_external)})
    return result


def clear_health_cache() -> None:
    with _CACHE_LOCK:
        _CACHE.update({"at": 0.0, "result": None, "external": False})
