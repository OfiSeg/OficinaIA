"""Migración inicial, de una sola vez, de los XLSX del repositorio a Google Sheets.

Requiere las mismas variables de entorno que la aplicación. No crea planillas:
los dos SHEET_ID deben existir y estar compartidos con la cuenta de servicio.
"""
from pathlib import Path

from dotenv import load_dotenv
from openpyxl import load_workbook

import google_sheets_service

BASE_DIR = Path(__file__).resolve().parent
LIBROS = {
    "1": BASE_DIR / "excel_interno.xlsx",
    "2": BASE_DIR / "excel_flotas.xlsx",
}

# El ZIP limpio de GitHub no trae excel_flotas.xlsx. Si no existe, dejamos
# igualmente inicializado el libro 2 con el esquema que usa /flota, para que
# el primer guardado no falle por ausencia de encabezados.
HEADERS_FLOTAS = [[
    "ITEM", "MARCA", "MODELO", "AÑO", "COBERTURA SOLICITADA", "USO",
    "MOTOR", "CHASIS", "ACCESORIO", "PATENTE", "SUMA ASEGURADA",
]]


def matriz_desde_xlsx(path: Path) -> tuple[str, list[list[str]]]:
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb.active
        filas = [
            ["" if value is None else str(value) for value in row]
            for row in ws.iter_rows(values_only=True)
        ]
        return ws.title, filas
    finally:
        wb.close()


def main() -> int:
    load_dotenv(BASE_DIR / ".env")
    migrados = 0
    for libro_id, path in LIBROS.items():
        if path.is_file():
            hoja, filas = matriz_desde_xlsx(path)
            google_sheets_service.guardar_excel(filas, hoja, libro_id)
            print(f"Libro {libro_id}: {len(filas)} filas migradas desde {path.name}.")
            migrados += 1
            continue

        if libro_id == "2":
            google_sheets_service.guardar_excel(HEADERS_FLOTAS, "Datos", libro_id)
            print("Libro 2: excel_flotas.xlsx no existe; se creó el esquema inicial de Flotas en Sheets.")
            migrados += 1
        else:
            print(f"Libro {libro_id}: {path.name} no existe; no se pudo migrar.")
    return 0 if migrados else 1


if __name__ == "__main__":
    raise SystemExit(main())
