"""Regresión determinística del generador de carta de cotización.

No usa red, IA ni datos reales de la oficina. Trabaja sólo con la plantilla y
assets versionados del proyecto dentro de un directorio temporal.
"""
from __future__ import annotations

from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory
import shutil
import zipfile

import fitz
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm

from cotizacion_document_service import (
    AZUL_INSTITUCIONAL,
    CotizacionDocumentService,
    ESTILOS_REQUERIDOS,
    MARCADOR,
    catalogo_companias,
    normalizar_porcentaje,
    normalizar_cobertura_para_carta,
)

BASE = Path(__file__).resolve().parent
TEMPLATE = BASE / "plantillas" / "word final.docx"
LOGOS = BASE / "static" / "img" / "companias"
CIAS = [
    "ATM",
    "Federación Patronal",
    "Mercantil Andina",
    "San Cristóbal",
    "Agrosalta",
    "Rivadavia",
    "Allianz",
]


def check(condicion, mensaje):
    if not condicion:
        raise AssertionError(mensaje)


def alternativa(compania: str, idx: int, *, larga=False):
    contenidos = [
        {"tipo": "beneficio", "texto": "Responsabilidad civil"},
        {"tipo": "beneficio", "texto": "Incendio total y parcial"},
        {"tipo": "beneficio", "texto": "Robo/hurto total y parcial"},
        {"tipo": "beneficio", "texto": "Destrucción total por accidente"},
    ]
    if larga:
        contenidos += [
            {"tipo": "beneficio", "texto": "Cristales laterales y cerraduras según condiciones de póliza"},
            {"tipo": "beneficio", "texto": "Daños parciales por accidente según la alternativa seleccionada"},
            {"tipo": "nota", "texto": "En caso de daño parcial se aplican los límites y condiciones informados en la cotización."},
        ]
    data = {
        "compania": compania,
        "nombre": "Todo Riesgo" if idx % 3 == 0 else "Terceros Completo",
        "suma": "$44.415.000",
        "franquicia_pct": "6" if idx % 3 == 0 else "",
        "franquicia_importe": "$2.664.900" if idx % 3 == 0 else "",
        "precio": f"${190 + idx}.000",
        "contenidos": contenidos,
    }
    if compania == "ATM":
        data.update(precio="", cuponera="$158.561", adherido="$79.280")
    return data


def texto_docx(doc: Document) -> str:
    partes = [p.text for p in doc.paragraphs]
    def recorrer_tabla(table):
        for row in table.rows:
            for cell in row.cells:
                partes.extend(p.text for p in cell.paragraphs)
                for nested in cell.tables:
                    recorrer_tabla(nested)
    for table in doc.tables:
        recorrer_tabla(table)
    return "\n".join(partes)


def validar_docx(path: Path, *, logos_esperados: int, baseline_shapes: int, companias_esperadas: int | None = None):
    doc = Document(path)
    text = texto_docx(doc)
    check(MARCADOR not in text, "El marcador {{COTIZACION}} quedó visible.")
    check("**" not in text, "Quedó Markdown visible en el Word.")
    low = text.lower()
    check("undefined" not in low and "null" not in low, "Aparecieron valores vacíos inválidos.")
    check("%%" not in text, "La franquicia duplicó el signo de porcentaje.")
    check("$ " not in text, "Quedó un espacio inconsistente después del signo $.")
    for estilo in ESTILOS_REQUERIDOS:
        check(estilo in doc.styles, f"Se perdió el estilo Word {estilo}.")
    check(len(doc.inline_shapes) == baseline_shapes + logos_esperados, "Cantidad inesperada de logos en el Word.")
    # Los logos comparten un slot óptico, pero cada asset puede tener una altura
    # recomendada propia para igualar peso visual sin estirarlo (p. ej. Allianz
    # y San Cristóbal). Nunca deben salirse de los límites del sistema.
    min_h, max_h = int(Cm(0.35)), int(Cm(0.75))
    for shape in list(doc.inline_shapes)[baseline_shapes:]:
        check(min_h <= int(shape.height) <= max_h, "Un logo quedó fuera del slot óptico permitido.")

    if companias_esperadas is not None:
        check(len(doc.tables) == companias_esperadas, "Cantidad inesperada de secciones/tablas de compañía.")
    for table in doc.tables:
        check(len(table.rows) >= 1, "Una sección de compañía quedó vacía.")
        first_row = table.rows[0]
        first_tr_pr = first_row._tr.trPr
        check(first_tr_pr is not None and first_tr_pr.find(qn("w:cantSplit")) is not None, "Logo + primera cobertura pueden separarse entre páginas.")
        check(first_tr_pr.find(qn("w:tblHeader")) is None, "El logo volvió a configurarse como encabezado repetible.")
        header_cell = first_row.cells[0]
        header_p = header_cell.paragraphs[0]
        p_pr = header_p._p.pPr
        p_bdr = p_pr.find(qn("w:pBdr")) if p_pr is not None else None
        bottom = p_bdr.find(qn("w:bottom")) if p_bdr is not None else None
        check(bottom is not None and bottom.get(qn("w:color")) == AZUL_INSTITUCIONAL, "Falta la línea azul institucional debajo del logo.")
        check(bottom.get(qn("w:sz")) in {"4", "5", "6"}, "La línea de compañía no es fina.")

        # La primera cobertura comparte la fila con el logo para que ambos se
        # muevan juntos; las siguientes tienen su propia fila indivisible.
        for row_index, row in enumerate(table.rows):
            nested = row.cells[0].tables
            check(len(nested) == 1, "Cada cobertura debe vivir en un único bloque documental interno.")
            tr_pr = nested[0].rows[0]._tr.trPr
            check(tr_pr is not None and tr_pr.find(qn("w:cantSplit")) is not None, "Una cobertura normal puede dividirse entre páginas.")
            if row_index < len(table.rows) - 1:
                tc_pr_nested = nested[0].rows[0].cells[0]._tc.tcPr
                borders_nested = tc_pr_nested.find(qn("w:tcBorders")) if tc_pr_nested is not None else None
                bottom_nested = borders_nested.find(qn("w:bottom")) if borders_nested is not None else None
                check(bottom_nested is not None and bottom_nested.get(qn("w:color")) == "DCE4EC", "Falta el separador suave entre coberturas.")
    cierres = [p for p in doc.paragraphs if p.text.strip() in {"Gracias por elegirnos", "Seguros San José"}]
    check(sum(p.text.strip() == "Gracias por elegirnos" for p in cierres) == 1, "El cierre institucional falta o está duplicado.")
    check(cierres[-2].text.strip() == "Gracias por elegirnos" and cierres[-1].text.strip() == "Seguros San José", "Las dos líneas del cierre no forman el último bloque.")
    for p in cierres[-2:]:
        frame = p._p.pPr.find(qn("w:framePr")) if p._p.pPr is not None else None
        check(frame is not None and frame.get(qn("w:yAlign")) == "bottom", "El cierre no quedó anclado al pie.")


def paginas_pdf(path: Path) -> int:
    with fitz.open(path) as pdf:
        paginas = [page.get_text() for page in pdf]
        texto = "\n".join(paginas)
        check(texto.count("Gracias por elegirnos") == 1, "El cierre debe aparecer una sola vez en el PDF.")
        check("Gracias por elegirnos" in paginas[-1] and all("Gracias por elegirnos" not in page for page in paginas[:-1]), "El cierre no quedó únicamente en la última página.")
        check(MARCADOR not in texto and "**" not in texto, "El PDF contiene marcadores/Markdown inválidos.")
        return len(pdf)


def paginas_texto_pdf(path: Path) -> list[str]:
    with fitz.open(path) as pdf:
        return [page.get_text() for page in pdf]


def main():
    check(TEMPLATE.exists(), "Falta plantillas/word final.docx")
    check((LOGOS / "logos_manifest.json").exists(), "Falta logos_manifest.json")
    service = CotizacionDocumentService(template_path=TEMPLATE, logos_dir=LOGOS)
    catalogo = catalogo_companias(LOGOS)
    # Las siete compañías con asset histórico conservan orden/identidad. El
    # catálogo universal puede sumar compañías sin logo para los selectores;
    # esa ampliación no debe invalidar la generación documental.
    con_logo = [x["key"] for x in catalogo if x.get("file")]
    check(con_logo[:7] == ["atm", "federacion_patronal", "mercantil_andina", "san_cristobal", "agrosalta", "rivadavia", "allianz"], "Las compañías documentales existentes perdieron orden o identidad.")
    check(normalizar_porcentaje("2%%") == "2%", "Falló la normalización pública de porcentaje.")
    base_doc = Document(TEMPLATE)
    baseline_shapes = len(base_doc.inline_shapes)

    for cia in CIAS:
        check(service.resolver_logo(cia) is not None, f"No se resolvió el logo de {cia}.")
    for alias in ("Federacion Patronal", "FEDERACION", "MERCANTIL", "San Cristobal", "GRUPO SAN CRISTOBAL", "Seguros Rivadavia", "Allianz Seguros"):
        check(service.resolver_logo(alias) is not None, f"No se resolvió el alias {alias}.")
    _, alto_sc = service._dimensiones_logo_cm("San Cristóbal", service.resolver_logo("San Cristóbal"))
    _, alto_allianz = service._dimensiones_logo_cm("Allianz", service.resolver_logo("Allianz"))
    _, alto_atm = service._dimensiones_logo_cm("ATM", service.resolver_logo("ATM"))
    check(alto_sc < alto_atm and alto_allianz < alto_atm, "Allianz/San Cristóbal no redujeron su peso visual respecto del slot estándar.")

    with TemporaryDirectory(prefix="oficinaia_validacion_cotizacion_") as raw:
        tmp = Path(raw)
        paginas_observadas = set()
        for n in (1, 2, 3, 5, 7):
            alts = [alternativa(CIAS[i % len(CIAS)], i) for i in range(n)]
            datos = {"vehiculo": "Peugeot 208 2026", "alternativas": alts}
            docx = service.generar_docx(datos, tmp / f"cotizacion_{n}.docx")
            validar_docx(docx, logos_esperados=n, baseline_shapes=baseline_shapes, companias_esperadas=n)
            if service._libreoffice_bin():
                pdf = service.convertir_pdf(docx, tmp / f"pdf_{n}")
                paginas_observadas.add(paginas_pdf(pdf))
                imagenes = service.generar_imagenes(pdf, tmp / f"png_{n}", "png")
                check(len(imagenes) == paginas_pdf(pdf), "PNG no coincide con cantidad de páginas PDF.")
                jpgs = service.generar_imagenes(pdf, tmp / f"jpg_{n}", "jpg")
                check(len(jpgs) == paginas_pdf(pdf), "JPG no coincide con cantidad de páginas PDF.")

        # Fuerza un caso largo para verificar flujo natural a tres o más páginas.
        alts_largas = [alternativa(CIAS[i % len(CIAS)], i, larga=True) for i in range(7)]
        long_docx = service.generar_docx({"vehiculo": "Prueba Larga", "alternativas": alts_largas}, tmp / "cotizacion_larga.docx")
        validar_docx(long_docx, logos_esperados=7, baseline_shapes=baseline_shapes, companias_esperadas=7)
        if service._libreoffice_bin():
            long_pdf = service.convertir_pdf(long_docx, tmp / "pdf_larga")
            long_pages = paginas_pdf(long_pdf)
            check(long_pages >= 3, f"La prueba larga debía fluir a 3+ páginas y produjo {long_pages}.")
            paginas_observadas.add(long_pages)


        # Varias coberturas de la misma compañía: logo una sola vez por grupo
        # canónico, incluso usando aliases, y normalización de % / moneda.
        repetidas = [
            alternativa("ATM", 0),
            alternativa("ATM", 1),
            alternativa("Mercantil Andina", 2),
            alternativa("MERCANTIL", 3),
            alternativa("ATM", 4),
        ]
        repetidas[0]["suma"] = "$ 69.000.000"
        repetidas[2]["franquicia_pct"] = "2%"
        repetidas_docx = service.generar_docx(
            {"vehiculo": "Prueba Agrupación", "alternativas": repetidas},
            tmp / "cotizacion_agrupada.docx",
        )
        validar_docx(repetidas_docx, logos_esperados=2, baseline_shapes=baseline_shapes, companias_esperadas=2)
        repetidas_text = texto_docx(Document(repetidas_docx))
        check("ATM" not in repetidas_text and "MERCANTIL ANDINA" not in repetidas_text, "El documento repite nombres de compañía aunque ya muestra sus logos.")
        check("Suma asegurada · $69.000.000" in repetidas_text, "No se normalizó el espacio monetario.")
        check("Franquicia · 2%" in repetidas_text and "2%%" not in repetidas_text, "No se normalizó la franquicia porcentual.")
        nota_2 = "En caso de daño parcial, queda a cargo del asegurado una franquicia equivalente al 2% de la suma asegurada. Todo gasto que supere ese importe queda a cargo de la compañía."
        check(nota_2 in repetidas_text, "La explicación universal de franquicia no llegó al documento.")
        check(repetidas_text.count(nota_2) == 1, "La explicación de franquicia quedó duplicada en el documento.")

        # Idempotencia: validar un payload ya normalizado no puede volver a
        # agregar una nota, variante o beneficio.
        tr_payload = {"vehiculo": "Idempotencia", "alternativas": [{
            "compania": "ATM", "familia": "TODO_RIESGO", "nombre": "Todo Riesgo",
            "nombre_comercial": "Todo Riesgo", "franquicia_pct": "3",
            "contenidos": [
                {"tipo": "beneficio", "texto": "Responsabilidad civil, incendio total y parcial, robo/hurto total y parcial, destrucción total y daños parciales por accidente"},
                {"tipo": "nota", "texto": "En caso de daño parcial, queda a cargo del asegurado una franquicia equivalente al 3% de la suma asegurada. Todo gasto que supere ese importe queda a cargo de la compañía."},
            ],
        }]}
        una = service.validar_datos(tr_payload)
        dos = service.validar_datos(una)
        check(una == dos, "La normalización documental no es idempotente.")
        check(sum(1 for x in dos["alternativas"][0]["contenidos"] if x.get("tipo") == "nota") == 1, "Una segunda validación duplicó la franquicia.")

        # C_PLUS/Full hereda siempre el núcleo de Terceros Completo y suma sólo
        # adicionales confirmados.
        plus = service.validar_datos({"vehiculo": "Plus", "alternativas": [{
            "compania": "Mercantil Andina", "familia": "C_PLUS",
            "nombre": "Terceros Completo M Plus", "nombre_comercial": "Terceros Completo Plus",
            "tiene_grua": True,
            "riesgos_detectados": ["RESPONSABILIDAD_CIVIL", "INCENDIO_TOTAL", "INCENDIO_PARCIAL", "ROBO_HURTO_TOTAL", "ROBO_HURTO_PARCIAL", "DESTRUCCION_TOTAL_ACCIDENTE", "RUEDAS", "VIDRIOS", "GRANIZO", "CERRADURAS", "GRUA"],
            "contenidos": [
                {"tipo": "beneficio", "texto": "Ruedas, vidrios, granizo y cerraduras"},
                {"tipo": "beneficio", "texto": "Incluye grúa"},
            ],
        }]})["alternativas"][0]
        plus_items = [x["texto"] for x in plus["contenidos"] if x.get("tipo") != "nota"]
        for esperado in ["Responsabilidad Civil", "Incendio Total y Parcial", "Robo/Hurto Total y Parcial", "Destrucción Total por Accidente"]:
            check(esperado in plus_items, f"C_PLUS perdió la prestación canónica: {esperado}.")
        for esperado in ["Ruedas", "Vidrios", "Granizo", "Cerraduras", "Incluye grúa"]:
            check(esperado in plus_items, f"C_PLUS perdió el adicional confirmado: {esperado}.")
        check(not any(", incendio" in x.lower() and ", robo" in x.lower() for x in plus_items), "C_PLUS volvió a usar una frase legacy en lugar de una prestación por bullet.")


        # Toda cobertura usa prestaciones canónicas, una por bullet. C1 sin DT
        # no vuelve al speech legacy; C con DT agrega sólo el bullet correspondiente.
        c1 = service.validar_datos({"vehiculo":"AgroSalta", "alternativas":[{
            "compania":"AgroSalta", "codigo":"C1", "familia":"C1",
            "nombre":"Terceros Completo", "nombre_comercial":"Terceros Completo",
            "suma":"$8.999.999", "precio":"$72.000",
            "riesgos_detectados":["RESPONSABILIDAD_CIVIL","INCENDIO_TOTAL","INCENDIO_PARCIAL","ROBO_HURTO_TOTAL","ROBO_HURTO_PARCIAL"],
            "contenidos":[{"tipo":"beneficio","texto":"Responsabilidad civil, incendio total y parcial y robo/hurto total y parcial"}],
        }]})["alternativas"][0]
        check([x["texto"] for x in c1["contenidos"] if x["tipo"] == "beneficio"] == [
            "Responsabilidad Civil", "Incendio Total y Parcial", "Robo/Hurto Total y Parcial"
        ], "C1 no quedó expresado con prestaciones canónicas por bullet.")
        c = service.validar_datos({"vehiculo":"AgroSalta", "alternativas":[{
            "compania":"AgroSalta", "codigo":"C", "familia":"C",
            "nombre":"Terceros Completo", "nombre_comercial":"Terceros Completo",
            "suma":"$8.999.999", "precio":"$73.000",
            "riesgos_detectados":["RESPONSABILIDAD_CIVIL","INCENDIO_TOTAL","INCENDIO_PARCIAL","ROBO_HURTO_TOTAL","ROBO_HURTO_PARCIAL","DESTRUCCION_TOTAL_ACCIDENTE"],
            "contenidos":[],
        }]})["alternativas"][0]
        check([x["texto"] for x in c["contenidos"] if x["tipo"] == "beneficio"] == [
            "Responsabilidad Civil", "Incendio Total y Parcial", "Robo/Hurto Total y Parcial", "Destrucción Total por Accidente"
        ], "C con DT no quedó expresado con prestaciones canónicas por bullet.")

        tech = service.validar_datos({"vehiculo":"ATM", "alternativas":[{
            "compania":"ATM", "codigo":"CPr", "familia":"C_PLUS",
            "nombre":"Terceros Completo Premium", "nombre_comercial":"Terceros Completo Premium",
            "riesgos_detectados":["RESPONSABILIDAD_CIVIL","INCENDIO_TOTAL","INCENDIO_PARCIAL","ROBO_HURTO_TOTAL","ROBO_HURTO_PARCIAL","DESTRUCCION_TOTAL_ACCIDENTE"],
            "detalle_tecnico":["Granizo hasta la suma asegurada · 2 eventos por año"],
            "contenidos":[],
        }]})["alternativas"][0]
        check(tech.get("detalle_tecnico") == ["Granizo hasta la suma asegurada · 2 eventos por año"], "Se perdió la metadata técnica de una variante ATM.")

        for codigo, nombre in (("C", "Terceros Completo Plus"), ("CPr", "Terceros Completo Premium"), ("CB", "Terceros Completo Black")):
            variante_atm = normalizar_cobertura_para_carta({
                "compania": "ATM", "codigo": codigo, "familia": "C_PLUS",
                "nombre": nombre, "nombre_comercial": nombre,
                "riesgos_detectados": ["RESPONSABILIDAD_CIVIL", "INCENDIO_TOTAL", "INCENDIO_PARCIAL", "ROBO_HURTO_TOTAL", "ROBO_HURTO_PARCIAL", "DESTRUCCION_TOTAL_ACCIDENTE"],
                "contenidos": [],
            })
            check(variante_atm["nombre"].upper() == nombre.upper(), f"ATM {codigo} perdió su variante comercial y fue aplastada por C_PLUS.")

        atm_b_estructurado = normalizar_cobertura_para_carta({
            "compania": "ATM", "codigo": "B", "familia": "B",
            "nombre": "Robo e Incendio Total y/o Parcial + Accidente Total",
            "nombre_comercial": "Robo e Incendio Total y/o Parcial + Accidente Total",
            "riesgos_detectados": ["RESPONSABILIDAD_CIVIL", "INCENDIO_TOTAL", "INCENDIO_PARCIAL", "ROBO_HURTO_TOTAL", "ROBO_HURTO_PARCIAL", "DESTRUCCION_TOTAL_ACCIDENTE"],
            "contenidos": [],
        })
        atm_b_items = [x["texto"] for x in atm_b_estructurado["contenidos"] if x["tipo"] == "beneficio"]
        check("Incendio Total y Parcial" in atm_b_items and "Robo/Hurto Total y Parcial" in atm_b_items, "La familia B volvió a pisar los riesgos estructurados ATM.")

        # Compañía desconocida: fallback textual y generación no bloqueada.
        unknown = {"vehiculo": "Prueba", "alternativas": [alternativa("Compañía Sin Logo", 1)]}
        unknown_docx = service.generar_docx(unknown, tmp / "sin_logo.docx")
        validar_docx(unknown_docx, logos_esperados=0, baseline_shapes=baseline_shapes, companias_esperadas=1)
        check("COMPAÑÍA SIN LOGO" in texto_docx(Document(unknown_docx)), "Falló el fallback textual de compañía.")

        # La confirmación manual prevalece sobre la detección y determina logo,
        # encabezado y agrupación documental.
        corregida = alternativa("Mercantil Andina", 1)
        corregida.update(compania_detectada="Mercantil Andina", compania_confirmada="AgroSalta")
        corregida_docx = service.generar_docx({"vehiculo": "Prueba", "alternativas": [corregida]}, tmp / "compania_corregida.docx")
        corregida_text = texto_docx(Document(corregida_docx))
        check("AGROSALTA" not in corregida_text and "MERCANTIL ANDINA" not in corregida_text, "El Word volvió a repetir el nombre de la compañía junto al logo.")
        # La prioridad de la compañía confirmada sigue validándose por resolución
        # de asset: AgroSalta debe ser el logo canónico del grupo corregido.
        check(service.resolver_logo("AgroSalta").name == "agrosalta.png", "La compañía confirmada no resolvió su logo canónico.")

        variante = alternativa("ATM", 0)
        variante.update(
            nombre="Responsabilidad Civil · Con grúa",
            variante_grua="CON GRÚA",
            franquicia_pct="",
            franquicia_importe="",
            contenidos=[{"tipo": "beneficio", "texto": "Responsabilidad Civil"}],
        )
        variante_docx = service.generar_docx({"vehiculo": "Prueba", "alternativas": [variante]}, tmp / "variante_grua.docx")
        variante_text = texto_docx(Document(variante_docx))
        check("RESPONSABILIDAD CIVIL · CON GRÚA" in variante_text, "La variante de grúa no quedó en el título.")

        # La carta normaliza códigos técnicos por riesgos reales y ordena de
        # básico a completo sin tocar el orden de compañías.
        tecnicas = [
            {
                "compania": "Mercantil Andina",
                "nombre": "Todo Riesgo",
                "suma": "$69.900.600",
                "franquicia_pct": "2%%",
                "franquicia_importe": "$1.398.012",
                "precio": "$247.000",
                "contenidos": [
                    {"tipo": "beneficio", "texto": "Responsabilidad Civil"},
                    {"tipo": "beneficio", "texto": "Incendio Total y Parcial"},
                    {"tipo": "beneficio", "texto": "Robo Total y Parcial"},
                    {"tipo": "beneficio", "texto": "Destrucción Total y daños parciales por accidente"},
                ],
            },
            {
                "compania": "Mercantil Andina",
                "nombre": "Cobertura B0",
                "suma": "$69.900.600",
                "precio": "$83.000",
                "contenidos": [
                    {"tipo": "beneficio", "texto": "Responsabilidad civil"},
                    {"tipo": "beneficio", "texto": "Robo Total"},
                ],
            },
            {
                "compania": "Mercantil Andina",
                "nombre": "Responsabilidad Civil",
                "suma": "$69.900.600",
                "precio": "$62.000",
                "contenidos": [{"tipo": "beneficio", "texto": "Responsabilidad civil"}],
            },
            {
                "compania": "Mercantil Andina",
                "nombre": "Cobertura B3",
                "suma": "$69.900.600",
                "precio": "$74.000",
                "contenidos": [
                    {"tipo": "beneficio", "texto": "Responsabilidad civil"},
                    {"tipo": "beneficio", "texto": "Incendio total y parcial"},
                ],
            },
        ]
        tecnicas_docx = service.generar_docx({"vehiculo": "Normalización", "alternativas": tecnicas}, tmp / "normalizacion.docx")
        validar_docx(tecnicas_docx, logos_esperados=1, baseline_shapes=baseline_shapes, companias_esperadas=1)
        tecnicas_text = texto_docx(Document(tecnicas_docx))
        check("COBERTURA B0" not in tecnicas_text and "COBERTURA B3" not in tecnicas_text, "Quedó nomenclatura técnica evitable como título.")
        check("Robo/Hurto Total" in tecnicas_text, "No se normalizó la terminología Robo/Hurto.")
        check("Daños Parciales por Accidente" in tecnicas_text, "No se separó correctamente el riesgo de daños parciales.")
        posiciones = [
            tecnicas_text.index("RESPONSABILIDAD CIVIL"),
            tecnicas_text.index("ROBO TOTAL"),
            tecnicas_text.index("INCENDIO TOTAL Y PARCIAL"),
            tecnicas_text.index("TODO RIESGO"),
        ]
        check(posiciones == sorted(posiciones), "Las coberturas no quedaron ordenadas de básica a completa.")

        # Allianz: D4 conserva el producto base; Granizo es una prestación/
        # estado separado y nunca parte de la identidad visible de la cobertura.
        d4 = normalizar_cobertura_para_carta({
            "nombre": "D4 ALTA GAMA VIP - 2% - C/GRANIZO",
            "franquicia_pct": "2%%",
            "contenidos": [{"tipo": "beneficio", "texto": "Granizo"}],
        })
        check(d4["nombre"] == "ALTA GAMA VIP", "D4 no quedó como producto base limpio; Granizo no debe formar parte del título.")

        # La autoridad comercial del cotizador debe sobrevivir intacta al DOCX/PDF.
        explicita = normalizar_cobertura_para_carta({
            "nombre": "CÓDIGO TÉCNICO X9",
            "familia": "C_PLUS",
            "nombre_comercial": "Terceros Completo Plus",
            "variante_comercial": "CON GRÚA",
            "variante_grua": "CON GRÚA",
            "contenidos": [{"tipo": "beneficio", "texto": "Responsabilidad Civil"}],
        })
        check(explicita["nombre"] == "TERCEROS COMPLETO PLUS · CON GRÚA", "El renderer documental contradijo el nombre/variante comercial del cotizador.")

        # Si llega un payload antiguo con nombre técnico, la familia normalizada
        # sigue siendo suficiente para evitar 'Cobertura LB/LB1' en documentos.
        lb = normalizar_cobertura_para_carta({
            "nombre": "Cobertura LB",
            "familia": "LB",
            "contenidos": [{"tipo": "beneficio", "texto": "Responsabilidad Civil"}],
        })
        check(lb["nombre_comercial"] == "ROBO E INCENDIO + ROBO PARCIAL AL AMPARO + ACCIDENTE TOTAL", "LB volvió a salir como código técnico en documentos.")

        # Compañía extensa: logo único, primera cobertura unida al encabezado
        # y ninguna cobertura puede quedar partida entre páginas.
        misma_cia = []
        for idx in range(1, 9):
            item = alternativa("ATM", idx)
            item["nombre"] = f"Cobertura B{idx}"
            item["cuponera"] = f"${100 + idx}.000"
            item["adherido"] = f"${80 + idx}.000"
            item["precio"] = ""
            item["contenidos"] = [
                {"tipo": "nota", "texto": f"ID COBERTURA {idx}"},
                {"tipo": "beneficio", "texto": "Responsabilidad Civil"},
                {"tipo": "beneficio", "texto": "Incendio Total y Parcial"},
                {"tipo": "beneficio", "texto": "Robo Total y Parcial"},
                {"tipo": "beneficio", "texto": "Destrucción Total por Accidente"},
                {"tipo": "beneficio", "texto": "Ruedas, vidrios, granizo, cerraduras y grúa"},
            ]
            misma_cia.append(item)
        misma_docx = service.generar_docx({"vehiculo": "Continuidad", "alternativas": misma_cia}, tmp / "misma_compania.docx")
        validar_docx(misma_docx, logos_esperados=1, baseline_shapes=baseline_shapes, companias_esperadas=1)
        if service._libreoffice_bin():
            misma_pdf = service.convertir_pdf(misma_docx, tmp / "pdf_misma")
            pages = paginas_texto_pdf(misma_pdf)
            check(len(pages) >= 2, "La prueba de continuación no llegó a varias páginas.")
            check("ID COBERTURA 1" in pages[0], "Una compañía extensa dejó la primera página vacía.")
            for page in pages:
                check("ATM" not in page, "El documento volvió a imprimir el nombre de la compañía junto al logo.")
            for idx in range(1, 9):
                id_txt = f"ID COBERTURA {idx}"
                precio = f"Precio adherido · ${80 + idx}.000"
                paginas_id = [i for i, page in enumerate(pages) if id_txt in page]
                paginas_precio = [i for i, page in enumerate(pages) if precio in page]
                check(len(paginas_id) == 1 and paginas_id == paginas_precio, f"La cobertura {idx} quedó partida entre páginas.")
            check("Ruedas, vidrios, granizo y cerraduras" in "\n".join(pages) and "Incluye grúa" in "\n".join(pages), "La normalización de beneficios perdió información real.")

        # Nombres de archivo y sanitización.
        nombre = service.nombre_base({"vehiculo": "Peugeot / 208"}, date(2026, 9, 13))
        check(nombre == "Cotizacion_Peugeot_208_13-09-2026", "Nombre de archivo inesperado.")

    print("OK - cotización documental: secciones de compañía, normalización comercial, bloques indivisibles, continuidad multipágina y exports validados")


if __name__ == "__main__":
    main()
